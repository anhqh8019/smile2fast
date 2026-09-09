from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class ExcelService:

    HEADERS = [
        "Sel",
        "Pst",
        "Folio #",
        "BC",
        "Voucher",
        "Code",
        "Description",
        "Rm",
        "Amount",
        "Origin Ar",
        "Ex.Rate",
        "GLAmount",
        "Ref#",
        "Comment",
        "Seri",
        "Bill#",
        "Csh",
        "Post time",
        "Bill ID",
        "Print Time",
        "ChrPOS",
    ]

    def create_smile_excel(self, rows: list[dict]) -> BytesIO:

        workbook = Workbook()

        sheet = workbook.active
        sheet.title = "Smile Transactions"

        # ==============================
        # HEADER
        # ==============================

        for col_index, header in enumerate(
            self.HEADERS,
            start=1
        ):
            cell = sheet.cell(
                row=1,
                column=col_index,
                value=header
            )

            cell.font = Font(
                bold=True
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="D9EAF7"
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

        # ==============================
        # DATA
        # ==============================

        for row_index, row in enumerate(
            rows,
            start=2
        ):

            values = [
                row.get("sel"),
                row.get("pst"),
                row.get("folio"),
                row.get("bc"),
                row.get("voucher"),
                row.get("code"),
                row.get("description"),
                row.get("room"),
                row.get("amount"),
                row.get("origin_amount"),
                row.get("exchange_rate"),
                row.get("gl_amount"),
                row.get("ref_no"),
                row.get("comment"),
                row.get("seri"),
                row.get("bill_no"),
                row.get("cashier"),
                row.get("post_time"),
                row.get("bill_id"),
                row.get("print_time"),
                row.get("chr_pos"),
            ]

            for col_index, value in enumerate(
                values,
                start=1
            ):
                sheet.cell(
                    row=row_index,
                    column=col_index,
                    value=value
                )

        # ==============================
        # NUMBER FORMAT
        # ==============================

        for row_index in range(
            2,
            sheet.max_row + 1
        ):
            sheet.cell(
                row=row_index,
                column=9
            ).number_format = '#,##0.00'

            sheet.cell(
                row=row_index,
                column=10
            ).number_format = '#,##0.00'

            sheet.cell(
                row=row_index,
                column=11
            ).number_format = '0.00'

            sheet.cell(
                row=row_index,
                column=12
            ).number_format = '#,##0.00'

        # ==============================
        # WIDTH
        # ==============================

        widths = {
            1: 5,
            2: 5,
            3: 12,
            4: 6,
            5: 22,
            6: 8,
            7: 20,
            8: 10,
            9: 16,
            10: 16,
            11: 10,
            12: 16,
            13: 15,
            14: 15,
            15: 15,
            16: 12,
            17: 10,
            18: 25,
            19: 12,
            20: 20,
            21: 12,
        }

        for index, width in widths.items():

            sheet.column_dimensions[
                get_column_letter(index)
            ].width = width

        sheet.freeze_panes = "A2"

        # ==============================
        # OUTPUT
        # ==============================

        output = BytesIO()

        workbook.save(output)

        output.seek(0)

        return output